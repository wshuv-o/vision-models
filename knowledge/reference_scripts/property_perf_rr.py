"""
KBS "Property Performance Report -> Rent Roll"  ->  Excel
(Example Plaza style; a different report from the ALT_CMROLL_1 one.)

Colab:
    !pip install pdfplumber openpyxl pandas -q

Layout notes this parser has to cope with, all of them multi-line:
  * NAICS description wraps over up to 4 lines ("Professional," / "Scientific,
    and" / "Technical" / "Services").
  * Tenant name wraps over up to 3 lines ("Smith Jones" / "Brown &" / "Partners").
  * Each rent step occupies TWO physical lines: date + monthly rent + annual
    psf on one, the monthly PSF alone on the next.
  * Tenant groups ("Acme Bank", "Example Law Firm, P.A.") head a set of suites and
    are followed by a "Total for <group>" line.
  * Section headers sit at x=38, group headers at x=42 -- that 4pt difference
    is what separates them.

Totals are excluded: "Total for <group>", "<section> Totals", and the final
"Total" line.

Anything the parser had to disentangle is written with a YELLOW cell and a note
in the "Review Note" column.
"""
import re
import pdfplumber
import pandas as pd

PDF_PATH = r"C:\path\to\your\folder\rentrolls\Example Plaza  Rent Roll 08.26.pdf"
XLSX_PATH = r"C:\path\to\your\folder\rentrolls\Example Plaza Rent Roll 08.26.xlsx"

BODY_TOP = 132.0
SECTION_X0 = (34.0, 40.0)    # section headers / section totals
GROUP_X0 = (40.0, 48.0)      # group headers / group totals / NAICS text

TEXT_ZONES = {                # x0 ranges for the left-hand text columns
    "NAICS Descr": (40, 100),
    "Bldg ID":     (100, 140),
    "Suite #":     (140, 172),
    "Tenant Name": (172, 240),
}
# right-edge (x1) bins for every numeric / date column
NUM_BINS = [
    ("Leased SF",                   250, 295),
    ("Remeas SF",                   300, 332),
    ("% of Bldg",                   332, 372),
    ("_LeaseDate",                  372, 418),   # start on line 1, end on line 2
    ("_CurrentRent",                418, 470),   # rent on line 1, PSF on line 2
    ("Current Annual Rent (psf)",   470, 508),
    ("Step Date",                   508, 548),
    ("_StepRent",                   548, 598),   # rent with a date, else PSF
    ("Step Annual Rent (psf)",      598, 640),
    ("Reimb Method",                640, 670),
    ("Base Year Exp/Stop",          700, 730),
    ("Security Deposit",            740, 790),
    ("Unpaid/Outstanding TIs & LCs", 790, 845),
    ("Renewal Options Term/Rate",   845, 885),
    ("No. Amend",                   885, 918),
    ("Comments",                    918, 975),
]

BLDG_RE = re.compile(r"^[A-Z]{1,4}\d{3,6}$")
DATE_RE = re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$")
SECTIONS = ("Future", "Vacant Not Leased", "Leased and Occupied")
ADDL_RE = re.compile(r"^\.*Add'l$")

COLS = ["Suite #", "Bldg ID", "Tenant Name", "Group", "Section", "Row Kind",
        "NAICS Descr", "Leased SF", "Remeas SF", "% of Bldg",
        "Lease Start", "Lease End",
        "Current Monthly Base Rent", "Current Monthly PSF",
        "Current Annual Rent (psf)",
        "Step Date", "Step Monthly Base Rent", "Step Monthly PSF",
        "Step Annual Rent (psf)",
        "Reimb Method", "Base Year Exp/Stop", "Security Deposit",
        "Unpaid/Outstanding TIs & LCs", "Renewal Options Term/Rate",
        "No. Amend", "Comments", "Page", "Review Note"]

MONEYISH = re.compile(r"^\$?-?\(?[\d,]+\.?\d*\)?%?$")


def _num(t):
    t = t.strip().replace("$", "").replace(",", "").replace("%", "")
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def _rows(words, tol=2.0):
    out, cur = [], []
    for w in sorted(words, key=lambda a: a["top"]):
        if cur and w["top"] - cur[-1]["top"] > tol:
            out.append(sorted(cur, key=lambda a: a["x0"]))
            cur = []
        cur.append(w)
    if cur:
        out.append(sorted(cur, key=lambda a: a["x0"]))
    return out


def _zone(row, name):
    lo, hi = TEXT_ZONES[name]
    return [w for w in row if lo <= w["x0"] < hi]


def _bins(row):
    """Map bin-name -> raw text for this physical line."""
    got = {}
    for w in row:
        for col, lo, hi in NUM_BINS:
            if lo < w["x1"] <= hi:
                got.setdefault(col, w["text"])
                break
    return got


def parse(pdf_path=PDF_PATH):
    out, flags = [], {}

    def flag(idx, col, note):
        flags.setdefault(idx, {})
        prev = flags[idx].get(col)
        flags[idx][col] = f"{prev}; {note}" if prev else note

    with pdfplumber.open(pdf_path) as pdf:
        section = group = None
        cur = None          # index in `out` of the detail row being built
        naics, tenant = [], []
        last_step = None    # index of the row whose Step PSF is still pending

        def finish():
            """Write the accumulated wrapped text back onto the detail row."""
            nonlocal cur, naics, tenant
            if cur is None:
                return
            if naics:
                out[cur]["NAICS Descr"] = " ".join(naics)
            if tenant:
                name = " ".join(tenant)
                if ADDL_RE.match(tenant[0]) or name.startswith("...Add'l"):
                    name = re.sub(r"^\.*Add'l\s+Space:\s*", "", name).strip()
                    out[cur]["Row Kind"] = "Additional Space"
                out[cur]["Tenant Name"] = name or None
            cur, naics, tenant = None, [], []

        for pageno, page in enumerate(pdf.pages, 1):
            words = page.extract_words(x_tolerance=1.6, use_text_flow=True)
            for row in _rows(words):
                if row[0]["top"] < BODY_TOP or row[0]["top"] > 715:
                    continue
                x0 = row[0]["x0"]
                text = " ".join(w["text"] for w in row).strip()
                b = _bins(row)
                has_bldg = any(TEXT_ZONES["Bldg ID"][0] <= w["x0"] < TEXT_ZONES["Bldg ID"][1]
                               and BLDG_RE.match(w["text"]) for w in row)

                # ---- section header / section total -------------------------
                if SECTION_X0[0] <= x0 < SECTION_X0[1]:
                    finish()
                    last_step = None
                    if text.endswith("Totals"):
                        group = None
                    else:
                        section = text
                        group = None
                    continue

                # ---- detail line -------------------------------------------
                if has_bldg:
                    finish()
                    r = {c: None for c in COLS}
                    r.update({"Section": section, "Group": group,
                              "Row Kind": "Detail", "Page": pageno})
                    r["Bldg ID"] = " ".join(w["text"] for w in _zone(row, "Bldg ID"))
                    r["Suite #"] = " ".join(w["text"] for w in _zone(row, "Suite #")) or None
                    for col in ("Leased SF", "Remeas SF", "% of Bldg",
                                "Current Annual Rent (psf)", "Step Annual Rent (psf)",
                                "Base Year Exp/Stop", "Security Deposit",
                                "Unpaid/Outstanding TIs & LCs",
                                "Renewal Options Term/Rate", "No. Amend"):
                        if col in b:
                            r[col] = _num(b[col])
                    for col in ("Reimb Method", "Comments"):
                        if col in b:
                            r[col] = b[col]
                    r["Lease Start"] = b.get("_LeaseDate")
                    r["Current Monthly Base Rent"] = _num(b["_CurrentRent"]) if "_CurrentRent" in b else None
                    r["Step Date"] = b.get("Step Date")
                    if "_StepRent" in b:
                        r["Step Monthly Base Rent"] = _num(b["_StepRent"])
                    out.append(r)
                    cur = len(out) - 1
                    last_step = cur
                    naics = [w["text"] for w in _zone(row, "NAICS Descr")]
                    tenant = [w["text"] for w in _zone(row, "Tenant Name")]
                    continue

                # ---- group total / grand total ------------------------------
                if GROUP_X0[0] <= x0 < GROUP_X0[1] and text.startswith("Total"):
                    finish()
                    last_step = None
                    group = None
                    continue

                # ---- group header -------------------------------------------
                if GROUP_X0[0] <= x0 < GROUP_X0[1]:
                    # A group header runs continuously from x=42 and therefore
                    # spills into the gap between the NAICS column (ends ~100)
                    # and the tenant column (starts ~175). A wrapped
                    # NAICS/tenant continuation line straddles that gap with
                    # nothing in it, e.g. "Technical" + "Firm," + "P.A.".
                    wide = any(100 <= w["x0"] < 172 for w in row)
                    if not any(w["x0"] >= 240 for w in row) and (cur is None or wide):
                        finish()
                        last_step = None
                        group = text
                        continue

                # ---- continuation lines within a detail block ---------------
                if cur is None:
                    continue
                naics += [w["text"] for w in _zone(row, "NAICS Descr")]
                tenant += [w["text"] for w in _zone(row, "Tenant Name")]

                if "_LeaseDate" in b and out[cur]["Lease End"] is None:
                    out[cur]["Lease End"] = b["_LeaseDate"]
                if "_CurrentRent" in b and out[cur]["Current Monthly PSF"] is None:
                    out[cur]["Current Monthly PSF"] = _num(b["_CurrentRent"])

                if "Step Date" in b:
                    # a new rent step -> its own row, inheriting the suite
                    s = {c: None for c in COLS}
                    s.update({"Suite #": out[cur]["Suite #"], "Bldg ID": out[cur]["Bldg ID"],
                              "Section": section, "Group": group,
                              "Row Kind": "Rent Step", "Page": pageno,
                              "Step Date": b["Step Date"]})
                    if "_StepRent" in b:
                        s["Step Monthly Base Rent"] = _num(b["_StepRent"])
                    if "Step Annual Rent (psf)" in b:
                        s["Step Annual Rent (psf)"] = _num(b["Step Annual Rent (psf)"])
                    out.append(s)
                    last_step = len(out) - 1
                elif "_StepRent" in b and last_step is not None:
                    # PSF line belonging to the step directly above
                    if out[last_step]["Step Monthly PSF"] is None:
                        out[last_step]["Step Monthly PSF"] = _num(b["_StepRent"])
                    else:
                        flag(last_step, "Step Monthly PSF",
                             "more than one PSF line found for this step")
                elif "_StepRent" in b:
                    flag(cur, "Step Monthly PSF",
                         "step PSF value with no step above it")
        finish()

    df = pd.DataFrame(out, columns=COLS)

    # ---- post-hoc sanity flags ---------------------------------------------
    for i, r in df.iterrows():
        if r["Row Kind"] == "Detail":
            if r["Section"] != "Vacant Not Leased":
                if not r["Lease Start"]:
                    flags.setdefault(i, {})["Lease Start"] = "no lease start date found"
                if not r["Tenant Name"]:
                    flags.setdefault(i, {})["Tenant Name"] = "no tenant name found"
        if r["Row Kind"] == "Rent Step" and r["Step Monthly PSF"] is None:
            flags.setdefault(i, {})["Step Monthly PSF"] = "step has no PSF line beneath it"
        # a step's rent / psf should agree with each other via the suite's SF
        if r["Step Monthly Base Rent"] and r["Step Monthly PSF"]:
            sf = df.loc[i, "Leased SF"]
            if not sf:
                prev = df.loc[:i][df.loc[:i]["Row Kind"] == "Detail"]
                sf = prev.iloc[-1]["Leased SF"] if len(prev) else None
            if sf:
                implied = r["Step Monthly Base Rent"] / sf
                if abs(implied - r["Step Monthly PSF"]) > 0.02:
                    flags.setdefault(i, {})["Step Monthly PSF"] = (
                        f"step rent / {sf:,.0f} sf = {implied:,.2f}, "
                        f"but the PSF line reads {r['Step Monthly PSF']:,.2f}")

    for idx, cols in flags.items():
        df.at[idx, "Review Note"] = " | ".join(dict.fromkeys(cols.values()))
    return df, flags


def build_workbook(df, flags, xlsx_path=XLSX_PATH):
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    YELLOW = PatternFill("solid", fgColor="FFFF00")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name="Rent Roll", index=False)
        ws = xl.book["Rent Roll"]
        hdr = [c.value for c in ws[1]]
        col_of = {n: i + 1 for i, n in enumerate(hdr)}

        money = '#,##0.00;(#,##0.00);-'
        fmts = {"Leased SF": '#,##0;(#,##0);-', "Remeas SF": '#,##0;(#,##0);-',
                "% of Bldg": '0.00"%"',
                "Current Monthly Base Rent": money, "Current Monthly PSF": '#,##0.00',
                "Current Annual Rent (psf)": '#,##0.00',
                "Step Monthly Base Rent": money, "Step Monthly PSF": '#,##0.00',
                "Step Annual Rent (psf)": '#,##0.00',
                "Base Year Exp/Stop": money, "Security Deposit": money,
                "Unpaid/Outstanding TIs & LCs": money,
                "Renewal Options Term/Rate": '#,##0.00'}
        widths = {"Suite #": 9, "Bldg ID": 10, "Tenant Name": 32, "Group": 28,
                  "Section": 18, "Row Kind": 16, "NAICS Descr": 34,
                  "Lease Start": 11, "Lease End": 11, "Step Date": 11,
                  "Review Note": 62, "Page": 6}

        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        for i, name in enumerate(hdr, start=1):
            L = get_column_letter(i)
            ws.column_dimensions[L].width = widths.get(name, 15)
            c = ws[f"{L}1"]
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            if name in fmts:
                for cell in ws[L][1:]:
                    cell.number_format = fmts[name]

        thin = Side(style="thin", color="B0B0B0")
        for r_i, kind in enumerate(df["Row Kind"], start=2):
            if kind != "Rent Step":
                for c_i in range(1, len(hdr) + 1):
                    ws.cell(row=r_i, column=c_i).border = Border(top=thin)

        for idx, cols in flags.items():
            for col in cols:
                if col in col_of:
                    ws.cell(row=idx + 2, column=col_of[col]).fill = YELLOW
            ws.cell(row=idx + 2, column=col_of["Review Note"]).fill = YELLOW
        for cell in ws[get_column_letter(col_of["Review Note"])][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return xlsx_path


if __name__ == "__main__":
    df, flags = parse()
    build_workbook(df, flags)
    det = df[df["Row Kind"].isin(["Detail", "Additional Space"])]
    print(f"{len(df)} rows ({len(det)} suites + {len(df) - len(det)} rent-step lines)")
    print(f"yellow cells: {sum(len(v) for v in flags.values())} across {len(flags)} rows\n")
    for sec, g in det.groupby("Section"):
        print(f"  {sec:22s} {len(g):3d} suites  "
              f"SF={g['Leased SF'].sum():>9,.0f}  "
              f"rent={g['Current Monthly Base Rent'].sum():>12,.2f}")
    print("\n->", XLSX_PATH)
