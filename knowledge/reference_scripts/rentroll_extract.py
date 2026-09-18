"""
MRI "Rent Roll" (Report ID: ALT_CMROLL_1)  ->  Excel

Colab:
    !pip install pdfplumber openpyxl pandas -q
    # set PDF_PATH / XLSX_PATH below, then run.

Layout: one sheet, laid out like the printed report. Suite Id is the first
column. A suite's future rent increases (Cat / Date / Monthly Amount / PSF)
sit on the SAME sheet -- the first one inline on the suite's own row, the rest
on continuation rows directly beneath it. Suite Id / Space Id / Tenant Name
are repeated on those continuation rows so an increase can never be read
against the wrong suite, while the money and square-footage columns appear
only once per suite so column sums stay correct.

Totals are excluded: the per-suite "Total (Weighted Avg PSF)" lines, the
"Totals:" block and the "Grand Total:" block.

Anything the parser had to work to disentangle is written with a YELLOW cell
and an explanation in the "Review Note" column.

Numeric columns are right-aligned, so figures are binned on their RIGHT edge
(x1). Binning on the left edge breaks as soon as a number grows a digit.
"""
import re
import pdfplumber
import pandas as pd

PDF_PATH = r"C:\path\to\your\folder\rentrolls\Example Tower Rent Roll - 08.2026 (1).pdf"
XLSX_PATH = r"C:\path\to\your\folder\rentrolls\Example Tower Rent Roll - 08.2026.xlsx"

BODY_TOP = 130.0      # above this is the repeated page header
SUITE_X1 = 75.0       # suite id column ends here
NAME_X0_MAX = 300.0   # tenant name / dates live left of the numeric block
DATE_COL_X0 = 200.0   # dates normally start here; a name reaching this overlaps
INCR_X0 = 595.0       # right of this = "Future Rent Increases" block

NUM_BINS = [
    ("GLA SqFt",              300, 342),
    ("Monthly Base Rent",     342, 382),
    ("Annual Rate PSF",       382, 418),
    ("Monthly Cost Recovery", 418, 470),
    ("Security Deposit",      470, 510),
    ("Expense Stop",          510, 545),
    ("Monthly Other Income",  545, 595),
]
NUM_COLS = [c for c, _, _ in NUM_BINS]

SECTIONS = {"New Leases", "Vacant Suites", "Occupied Suites"}
DATE_RE = re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$")
TRAIL_DATE_RE = re.compile(r"(\d{1,2}/\d{1,2}/\d{2,4})$")
NUMISH_RE = re.compile(r"^\(?-?[\d,]+\.?\d*\)?$")
SUITEISH_RE = re.compile(r"^[A-Za-z0-9]+-[A-Za-z0-9]+$")

COLS = ["Suite Id", "Space Id", "Tenant Name", "Section", "Row Kind",
        "Rent Start", "Rent Expire"] + NUM_COLS + \
       ["Cat", "Date", "Monthly Amount", "PSF", "Property", "Page", "Review Note"]


def _num(t):
    t = t.replace(",", "").strip()
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def _rows(words, tol=2.0):
    out, cur = [], []
    for w in sorted(words, key=lambda a: a["top"]):
        if cur and w["top"] - cur[-1]["top"] > tol:
            out.append(sorted(cur, key=lambda a: a["x0"]))
            cur = []
        cur.append(w)
    if cur:
        out.append(sorted(cur, key=lambda a: a["x0"]))
    return out


def _property_name(pdf):
    for r in _rows(pdf.pages[0].extract_words(x_tolerance=1.6, use_text_flow=True)):
        if r[0]["top"] > BODY_TOP:
            break
        band = [w for w in r if 300 < w["x0"] < 600]
        txt = " ".join(w["text"] for w in band).strip()
        if txt and not txt.startswith(("Rent Roll", "Database", "Report")) \
                and not DATE_RE.match(txt):
            return txt
    return None


def _increase(right):
    cat = date = amt = psf = None
    for w in right:
        t, x1 = w["text"], w["x1"]
        if x1 <= 620:
            cat = t
        elif x1 <= 660:
            date = t
        elif x1 <= 740:
            amt = _num(t)
        else:
            psf = _num(t)
    return cat, date, amt, psf


def parse(pdf_path=PDF_PATH):
    """Return (DataFrame, flags) where flags maps row-index -> {column: note}."""
    out, flags = [], {}

    def flag(idx, col, note):
        flags.setdefault(idx, {})
        prev = flags[idx].get(col)
        flags[idx][col] = f"{prev}; {note}" if prev else note

    with pdfplumber.open(pdf_path) as pdf:
        prop = _property_name(pdf)
        section = None
        ctx = {"Suite Id": None, "Space Id": None, "Tenant Name": None}

        for pageno, page in enumerate(pdf.pages, 1):
            # use_text_flow stops an over-long tenant name that prints across
            # the date column from being interleaved character-by-character.
            words = page.extract_words(x_tolerance=1.6, use_text_flow=True)
            for r in _rows(words):
                if r[0]["top"] < BODY_TOP:
                    continue
                left = [w for w in r if w["x0"] < INCR_X0]
                right = [w for w in r if w["x0"] >= INCR_X0]
                ltext = " ".join(w["text"] for w in left).strip()

                if ltext in SECTIONS:
                    section = ltext
                    ctx = {"Suite Id": None, "Space Id": None, "Tenant Name": None}
                    continue
                # "Sqft:" only appears in the Totals / Grand Total block, whose
                # continuation lines do not start with the word "Total".
                if ltext.startswith(("Total", "Grand Total")) or "Sqft:" in ltext:
                    continue
                if not left and not right:
                    continue

                # ---- continuation line: one more future increase ------------
                if not left:
                    cat, date, amt, psf = _increase(right)
                    if not cat:
                        continue
                    row = {c: None for c in COLS}
                    row.update({"Suite Id": ctx["Suite Id"], "Space Id": ctx["Space Id"],
                                "Tenant Name": ctx["Tenant Name"], "Section": section,
                                "Row Kind": "Rent Increase", "Cat": cat, "Date": date,
                                "Monthly Amount": amt, "PSF": psf,
                                "Property": prop, "Page": pageno})
                    out.append(row)
                    if ctx["Suite Id"] is None:
                        flag(len(out) - 1, "Suite Id",
                             "increase line with no suite above it - check which suite it belongs to")
                    continue

                # ---- a suite / additional-space line ------------------------
                suite = " ".join(w["text"] for w in left if w["x1"] <= SUITE_X1).strip()
                mid = [w for w in left if w["x1"] > SUITE_X1 and w["x0"] < NAME_X0_MAX]

                dates, name_toks, notes = [], [], []
                for w in mid:
                    t = w["text"]
                    if DATE_RE.match(t):
                        dates.append(t)
                        continue
                    m = TRAIL_DATE_RE.search(t)
                    if m:  # name butted straight against its start date
                        head = t[:m.start()]
                        if head:
                            name_toks.append(dict(w, text=head))
                        dates.append(m.group(1))
                        notes.append(f"tenant name and date printed with no gap "
                                     f"({t!r}) - split into name + date")
                        continue
                    if w["x0"] >= DATE_COL_X0:
                        notes.append(f"tenant name runs into the date column ({t!r})")
                    name_toks.append(w)

                start = dates[0] if len(dates) > 0 else None
                expire = dates[1] if len(dates) > 1 else None
                if len(dates) > 2:
                    notes.append(f"{len(dates)} date-like values on this line")

                is_add = (len(name_toks) >= 2
                          and name_toks[0]["text"] == "Additional"
                          and name_toks[1]["text"] == "Space")
                space_id = None
                if is_add:
                    rest = name_toks[2:]
                    if rest and SUITEISH_RE.match(rest[0]["text"]):
                        space_id = rest[0]["text"]
                        rest = rest[1:]
                    name = " ".join(w["text"] for w in rest).strip() or ctx["Tenant Name"]
                    kind = "Additional Space"
                else:
                    name = " ".join(w["text"] for w in name_toks).strip() or None
                    kind = "Detail"

                if not suite and not is_add and not name:
                    continue

                row = {c: None for c in COLS}
                row.update({"Suite Id": suite or (ctx["Suite Id"] if is_add else None),
                            "Space Id": space_id, "Tenant Name": name,
                            "Section": section, "Row Kind": kind,
                            "Rent Start": start, "Rent Expire": expire,
                            "Property": prop, "Page": pageno})
                for w in left:
                    for col, lo, hi in NUM_BINS:
                        if lo < w["x1"] <= hi:
                            if row[col] is None:
                                row[col] = _num(w["text"])
                            break
                    else:
                        # a figure that fell outside every column bin would be
                        # silently lost -- never seen in this report, but flag
                        # it loudly rather than drop it.
                        if NUMISH_RE.match(w["text"]) and w["x1"] > 300:
                            notes.append(f"value {w['text']!r} did not land in any "
                                         f"known column (x1={w['x1']:.0f})")

                if right:
                    cat, date, amt, psf = _increase(right)
                    if cat:
                        row.update({"Cat": cat, "Date": date,
                                    "Monthly Amount": amt, "PSF": psf})

                out.append(row)
                idx = len(out) - 1

                for n in notes:
                    if "no gap" in n or "runs into the date column" in n:
                        flag(idx, "Tenant Name", n)
                        flag(idx, "Rent Start", n)
                    elif "did not land" in n:
                        flag(idx, "Row Kind", n)
                    else:
                        flag(idx, "Rent Start", n)
                if section != "Vacant Suites" and not start:
                    flag(idx, "Rent Start", "no rent start date found on this line")
                if section is None:
                    flag(idx, "Section", "row appeared before any section header")

                ctx = {"Suite Id": row["Suite Id"], "Space Id": space_id,
                       "Tenant Name": name}

    df = pd.DataFrame(out, columns=COLS)
    for idx, cols in flags.items():
        df.at[idx, "Review Note"] = " | ".join(dict.fromkeys(cols.values()))
    return df, flags


def build_workbook(df, flags, xlsx_path=XLSX_PATH):
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    YELLOW = PatternFill("solid", fgColor="FFFF00")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name="Rent Roll", index=False)
        ws = xl.book["Rent Roll"]

        money = '#,##0.00;(#,##0.00);-'
        fmts = {"GLA SqFt": '#,##0;(#,##0);-', "Monthly Base Rent": money,
                "Annual Rate PSF": '#,##0.00', "Monthly Cost Recovery": money,
                "Security Deposit": money, "Expense Stop": money,
                "Monthly Other Income": money, "Monthly Amount": money,
                "PSF": '#,##0.00'}
        widths = {"Suite Id": 14, "Space Id": 13, "Tenant Name": 36, "Section": 16,
                  "Row Kind": 16, "Rent Start": 11, "Rent Expire": 11, "Cat": 6,
                  "Date": 11, "Property": 12, "Page": 6, "Review Note": 60}

        hdr = [c.value for c in ws[1]]
        col_of = {name: i + 1 for i, name in enumerate(hdr)}
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions

        for i, name in enumerate(hdr, start=1):
            L = get_column_letter(i)
            ws.column_dimensions[L].width = widths.get(name, 15)
            c = ws[f"{L}1"]
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            if name in fmts:
                for cell in ws[L][1:]:
                    cell.number_format = fmts[name]

        # visually separate each suite block from the next
        thin = Side(style="thin", color="B0B0B0")
        for r_i, kind in enumerate(df["Row Kind"], start=2):
            if kind != "Rent Increase":
                for c_i in range(1, len(hdr) + 1):
                    ws.cell(row=r_i, column=c_i).border = Border(top=thin)

        # yellow = a cell the parser had to disentangle; eyeball it
        for idx, cols in flags.items():
            for col, note in cols.items():
                if col in col_of:
                    ws.cell(row=idx + 2, column=col_of[col]).fill = YELLOW
            ws.cell(row=idx + 2, column=col_of["Review Note"]).fill = YELLOW
        for cell in ws[get_column_letter(col_of["Review Note"])][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return xlsx_path


if __name__ == "__main__":
    df, flags = parse()
    build_workbook(df, flags)
    suites = df[df["Row Kind"] != "Rent Increase"]
    print(f"{len(df)} rows written "
          f"({len(suites)} suite/space lines + {len(df) - len(suites)} increase lines)")
    print(f"cells flagged yellow for review: {sum(len(v) for v in flags.values())} "
          f"across {len(flags)} rows")
    for sec in ["Occupied Suites", "Vacant Suites", "New Leases"]:
        s = suites[suites.Section == sec]
        print(f"  {sec:18s} {len(s):3d} lines  sqft={s['GLA SqFt'].sum():>10,.0f}  "
              f"base rent={s['Monthly Base Rent'].sum():>12,.2f}")
    print("\n->", XLSX_PATH)
